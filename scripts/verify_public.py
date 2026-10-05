"""Read-only checks of aggregate artifacts; does not replay statistical models."""
from __future__ import annotations

import ast
import csv
import hashlib
import json
from pathlib import Path
import re
import zipfile


def rows(path):
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle))


def verify_package(root):
    from openpyxl import load_workbook
    root = Path(root)
    checks = []
    def check(name, condition):
        checks.append({"check": name, "passed": bool(condition)})
    flow = rows(root / "tables/01_cohort_flow_and_roles_v32.csv")
    spec = json.loads((root / "models/model_specification_v32.json").read_text())
    expected = {"INSPIRE": (33394, 24872, 1680),
                "MOVER 2021": (2802, 2212, 274), "MOVER 2022": (2587, 2033, 259)}
    for cohort, values in expected.items():
        record = next(r for r in flow if r["cohort"] == cohort)
        check(cohort + "_denominators_events",
              tuple(int(float(record[k])) for k in
                    ["eligible_n", "observed_outcome_n", "observed_events"]) == values)
        model = spec["cohort_denominators"][cohort.replace(" ", "_")]
        check(cohort + "_model_denominators", (model["eligible"], model["observed"]) == values[:2])
    main = rows(root / "tables/36_main_table2_propagation_v32.csv")
    check("three_recalibration_strategies", len(main) == 3)
    for r, oe in zip(main, [1.487, 0.967, 1.012]):
        check("strategy_OE_" + r["update_estimator"], abs(float(r["oe_ratio"]) - oe) < 0.0005)
        check("strategy_AUROC_" + r["update_estimator"], abs(float(r["auroc"]) - 0.674) < 0.0005)
    canonical = {r["metric"]: r for r in rows(root / "tables/31_canonical_primary_bootstrap_v32.csv")}
    for metric in ["oe_ratio", "auroc", "brier", "calibration_slope"]:
        check("canonical_vs_primary_comparator_" + metric,
              abs(float(canonical[metric]["estimate"]) - float(main[2][metric])) < 1e-12)
    for key, column in [("alpha", "update_alpha"), ("beta", "update_beta")]:
        check("model_recalibration_" + key,
              abs(spec["recalibration"][key] - float(main[2][column])) < 1e-12)
    ties = rows(root / "audit/Table_S27_baseline_tie_sensitivity.csv")
    check("tie_retained_counts", [int(r["target_n"]) for r in ties] == [2587, 2587, 2586])
    check("tie_source_events", [int(r["source_events"]) for r in ties] == [1680, 1680, 1679])
    check("tie_target_observed_events", all(int(r["target_observed_events"]) == 259 for r in ties))
    tie_note = json.loads((root / "audit/baseline_tie_sensitivity.json").read_text())
    check("tie_no_bootstrap_no_clinical_adjudication",
          tie_note["bootstrap_rerun"] is False and tie_note["clinical_adjudication_performed"] is False)
    for record, other in zip(ties, tie_note["results"]):
        for key, value in other.items():
            check("tie_csv_json_" + record["scenario"] + "_" + key,
                  record[key] == value if isinstance(value, str) else abs(float(record[key]) - value) < 1e-12)
    book = load_workbook(root / "workbook/All_Revised_Tables.xlsx", data_only=True)
    formulas = load_workbook(root / "workbook/All_Revised_Tables.xlsx", data_only=False)
    check("Table_S27_present", {"Table_S27", "S27_numeric"} <= set(book.sheetnames))
    for sheet_name, file in [
        ("Primary_numeric", "tables/31_canonical_primary_bootstrap_v32.csv"),
        ("Comparators_numeric", "tables/36_main_table2_propagation_v32.csv"),
        ("S27_numeric", "audit/Table_S27_baseline_tie_sensitivity.csv"),
        ("S23_numeric", "tables/R1_surgical_flags.csv"),
        ("S24_numeric", "tables/32_stage_observation_balance_v32.csv"),
    ]:
        data = rows(root / file)
        sheet = book[sheet_name]
        headers = [c.value for c in sheet[1]]
        matching = True
        for index, record in enumerate(data, 2):
            for column, key in enumerate(headers, 1):
                actual, target = sheet.cell(index, column).value, record[key]
                if actual is None and target == "":
                    continue
                if isinstance(actual, bool):
                    same = str(actual).lower() == target.lower()
                else:
                    try:
                        same = abs(float(actual) - float(target)) < 1e-10
                    except (TypeError, ValueError):
                        same = str(actual) == target
                matching &= same
        check("workbook_csv_" + sheet_name, matching)
    coefficients = [book["Table_S6"].cell(i + 2, 3).value for i in range(len(spec["source_coefficients"]))]
    check("workbook_model_coefficients",
          all(abs(float(a) - b) < 5.1e-11 for a, b in zip(coefficients, spec["source_coefficients"])))
    check("workbook_missingness_suppressed", book["Table_S5"]["D13"].value == "<5"
          and book["Table_S5"]["E13"].value == "<0.18")
    check("workbook_no_cell_errors",
          not any(c.data_type == "e" for sheet in book for row in sheet for c in row))
    check("workbook_S27_formulas_retained",
          sum(c.data_type == "f" for row in formulas["Table_S27"] for c in row) == 3)
    unsafe = re.compile(r"/Users/|/Volumes/|/home/|file://|\bJAMIA\b|\bBJA\b", re.I)
    # Scan content, not the validator's own literal patterns.
    bad_docs = []
    for folder in ["config", "docs", "tables", "audit", "models"]:
        for path in (root / folder).glob("*"):
            if path.is_file() and path.suffix in {".json", ".csv", ".md", ".txt"}:
                if unsafe.search(path.read_text()):
                    bad_docs.append(path.relative_to(root).as_posix())
    check("no_private_paths_or_journal_naming_in_public_docs_data", not bad_docs)
    with zipfile.ZipFile(root / "workbook/All_Revised_Tables.xlsx") as archive:
        check("workbook_no_private_paths_or_journal_names",
              not any(unsafe.search(archive.read(n).decode("utf-8", "ignore"))
                      for n in archive.namelist() if n.endswith(".xml")))
        check("workbook_no_external_links", not any("externalLinks/" in n for n in archive.namelist()))
    forbidden = [p.relative_to(root).as_posix() for p in root.rglob("*")
                 if p.is_file() and (p.suffix in {".parquet", ".joblib", ".pkl", ".pickle",
                                                 ".docx", ".pyc", ".sqlite"}
                                    or "PRIVATE" in p.name or p.is_symlink())]
    check("no_restricted_artifact_types", not forbidden)
    check("no_obsolete_v31_output_files",
          not any("v31" in p.name for folder in ["tables", "models", "figures", "audit", "workbook"]
                  for p in (root / folder).glob("*")))
    source_map = json.loads((root / "docs/SOURCE_MAP.json").read_text())
    check("source_map_staged_hashes",
          all(hashlib.sha256((root / e["file"]).read_bytes()).hexdigest() == e["staged_sha256"]
              for e in source_map["files"]))
    for path in (root / "scripts").glob("*.py"):
        compile(path.read_text(), path.name, "exec")
    check("python_syntax", True)
    return {"scope": "Frozen aggregate comparison and packaging checks, not a statistical rerun.",
            "passed": all(c["passed"] for c in checks), "checks": checks,
            "workbook_sheets": len(book.sheetnames), "bad_public_text_files": bad_docs,
            "forbidden_artifacts": forbidden}
